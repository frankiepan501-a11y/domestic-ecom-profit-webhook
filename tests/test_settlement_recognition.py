import csv
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import parsers, settlement_engine, task_runner


def _csv_bytes(headers, rows):
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")

def workbook_bytes(title, headers, values):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    ws.append(values)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class SettlementRecognitionTests(unittest.TestCase):
    def test_tmall_ad_account_flow_excludes_prepayment_from_monthly_consumption(self):
        account_flow = (
            "记账时间,交易日期,收支类型,交易类型,操作金额(元),操作后余额(元),备注\n"
            "2026-07-01 02:00:00,2026-06-30,支出,扣款,8494.79,50.00,20260630现金消耗扣款\n"
            "2026-06-18 14:25:28,2026-06-18,支出,付款,897.04,10.00,下单金额被预支付扣款\n"
            "2026-06-18 14:25:28,2026-06-18,收入,充值,897.04,50.00,支付宝在线充值\n"
        ).encode("gbk")
        report = settlement_engine.SettlementReport("2026-06")

        _, ad_fee = settlement_engine.read_tmall_fees(
            report,
            [{"fname": "0947_fixture.csv", "buf": account_flow}],
            "天猫纷岚",
            {},
        )

        self.assertEqual(8494.79, ad_fee)
        self.assertEqual(8494.79, sum(row["spend"] for row in parsers.parse_tmall_ads(account_flow)))

    def test_douyin_empty_shop_still_records_ad_evidence_p0(self):
        result = settlement_engine.compute({"source_files": []}, {}, "2026-06")

        ad_gap_shops = {
            row[2]
            for row in result["gap_rows"]
            if row[0] == "P0" and row[1] == "抖音" and row[4:6] == ["资料缺口", "广告账单"]
        }
        self.assertEqual({"抖音宝空", "抖音纷岚"}, ad_gap_shops)

    def test_douyin_missing_ad_file_and_zero_spend_confirmation_is_p0(self):
        settlement = (
            "订单号,结算单类型,收入合计,结算金额,商品数量,商品ID,商品名称,平台服务费\n"
            "order-a,已结算,100,95,1,product-a,商品A,5\n"
        ).encode("utf-8-sig")
        orders = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,快递信息\n"
            "order-a,sub-a,商品A,product-a,SKU-A,1,SF-A-顺丰速运\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音结算订单.csv", "buf": settlement},
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音宝空订单明细.csv", "buf": orders},
            ],
            "logistics": [{"tracking": "SF-A", "carrier": "顺丰", "amount": 8, "source": "账单"}],
        }

        result = settlement_engine.compute(
            raw,
            {"SKU-A": {"unit_cost": 10, "source": "产品采购成本台", "name": "商品A"}},
            "2026-06",
        )

        ad_gaps = [
            row for row in result["gap_rows"]
            if row[0:6] == ["P0", "抖音", "抖音宝空", "2026-06", "资料缺口", "广告账单"]
        ]
        self.assertEqual(len(ad_gaps), 1)
        self.assertIn("抖音宝空订单明细.csv", ad_gaps[0][6])

    def test_douyin_confirmed_zero_ad_spend_is_recorded_and_not_a_gap(self):
        settlement = (
            "订单号,结算单类型,收入合计,结算金额,商品数量,商品ID,商品名称,平台服务费\n"
            "order-a,已结算,100,95,1,product-a,商品A,5\n"
        ).encode("utf-8-sig")
        orders = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,快递信息\n"
            "order-a,sub-a,商品A,product-a,SKU-A,1,SF-A-顺丰速运\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音结算订单.csv", "buf": settlement},
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音宝空订单明细.csv", "buf": orders},
                {"platform": "抖音", "shop": "纷岚店", "fname": "抖音纷岚结算订单.csv", "buf": settlement},
                {"platform": "抖音", "shop": "纷岚店", "fname": "抖音纷岚订单明细.csv", "buf": orders},
            ],
            "manifest_statuses": [
                {
                    "platform": "抖音",
                    "shop": "宝空店",
                    "file_type": "广告账单",
                    "status": "已确认无数据",
                },
                {
                    "platform": "抖音",
                    "shop": "纷岚店",
                    "file_type": "广告账单",
                    "status": "已确认无数据",
                },
            ],
            "logistics": [{"tracking": "SF-A", "carrier": "顺丰", "amount": 8, "source": "账单"}],
        }

        result = settlement_engine.compute(
            raw,
            {"SKU-A": {"unit_cost": 10, "source": "产品采购成本台", "name": "商品A"}},
            "2026-06",
        )

        for shop in ("抖音宝空", "抖音纷岚"):
            monthly = next(row for row in result["monthly_rows"] if row[1:3] == ["抖音", shop])
            self.assertEqual(0.0, monthly[10])
            self.assertFalse(any(
                row[1:6] == ["抖音", shop, "2026-06", "资料缺口", "广告账单"]
                for row in result["gap_rows"]
            ))
            self.assertTrue(any(
                row[0:6] == ["抖音", shop, "广告说明", "资料清单状态", 0, "已确认无数据"]
                for row in result["source_rows"]
            ))

    def test_douyin_order_csv_is_parsed_as_csv_not_excel(self):
        buf = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,商品金额,"
            "订单提交时间,支付完成时间,订单状态,售后状态,订单应付金额,快递信息\n"
            "P8001,S8001,手柄,10001,PK02-S2,1,99,2026-07-01 10:00:00,"
            "2026-07-01 10:01:00,已完成,-,99,SF123-顺丰速运,10001-1\n"
        ).encode("utf-8-sig")

        result = parsers.detect_and_parse("抖音宝空订单明细.csv", buf, "2026-07", "订单", platform="抖音")

        self.assertEqual("订单", result["kind"])
        self.assertEqual("P8001", result["data"][0]["main_oid"])
        self.assertEqual("PK02-S2", result["data"][0]["sku"])

    def test_xhs_plain_order_detail_filename_is_used_for_settlement_join(self):
        settlement = workbook_bytes(
            "商品结算明细",
            ["订单号", "交易类型", "商品数量", "商品实付/实退", "佣金总额", "商品名称"],
            ["P800347296683184211", "结算入账", 1, 99, -5, "手柄"],
        )
        order = workbook_bytes(
            "包裹详情",
            ["订单号", "商家编码", "商品名称", "SKU件数", "快递单号", "快递公司"],
            ["P800347296683184211", "PK02-S2", "手柄", 1, "SF123", "顺丰速运"],
        )
        raw = {
            "source_files": [
                {"platform": "小红书", "shop": "宝空店", "fname": "小红书宝空商品结算明细.xlsx", "buf": settlement},
                {"platform": "小红书", "shop": "宝空店", "fname": "小红书-宝空-2026-07-订单明细.xlsx", "buf": order},
            ],
            "logistics": [{"tracking": "SF123", "carrier": "顺丰", "amount": 10, "source": "API"}],
        }

        result = settlement_engine.compute(
            raw,
            {"PK02-S2": {"unit_cost": 20, "source": "产品采购成本台", "name": "ERP小红书手柄"}},
            "2026-07",
        )

        self.assertFalse(any("缺订单查询" in str(row[6]) for row in result["gap_rows"]))
        self.assertTrue(any(row[4] == "PK02-S2" for row in result["cost_rows"]))
        self.assertTrue(any(
            row[5] == "PK02-S2" and row[6] == "ERP小红书手柄"
            for row in result["product_rows"]
        ))

    def test_xhs_cancelled_order_without_waybill_is_recorded_as_no_freight(self):
        order_id = "P799684994323319861"
        settlement = workbook_bytes(
            "商品结算明细",
            ["订单号", "交易类型", "商品数量", "商品实付/实退", "佣金总额", "商品名称"],
            [order_id, "结算入账", 1, 214, -4.28, "手柄"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(settlement))
        ws = wb["商品结算明细"]
        ws.append([order_id, "退款", 1, -214, 4.28, "手柄"])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        settlement = buf.getvalue()
        order = workbook_bytes(
            "包裹详情",
            [
                "订单号", "订单状态", "售后状态", "商家编码", "商品名称",
                "SKU件数", "快递单号", "快递公司",
            ],
            [order_id, "已取消", "售后完成", "FF01A-04", "手柄", 1, "", ""],
        )
        raw = {
            "source_files": [
                {"platform": "小红书", "shop": "纷岚店", "fname": "小红书纷岚商品结算明细.xlsx", "buf": settlement},
                {"platform": "小红书", "shop": "纷岚店", "fname": "小红书纷岚订单明细.xlsx", "buf": order},
            ],
            "logistics": [],
        }

        result = settlement_engine.compute(
            raw,
            {"FF01A-04": {"unit_cost": 50, "source": "产品采购成本台", "name": "手柄"}},
            "2026-07",
        )

        logistics_gaps = [
            row for row in result["gap_rows"]
            if row[1] == "小红书" and row[4] == "物流成本" and row[5] == order_id
        ]
        self.assertEqual([], logistics_gaps)
        self.assertTrue(any(
            row[0] == "小红书" and row[3] == order_id and row[11] == "无需运费"
            for row in result["log_rows"]
        ))

    def test_xhs_cancelled_order_without_matching_refund_keeps_freight_p0(self):
        order_id = "P-CANCEL-NO-REFUND"
        settlement = workbook_bytes(
            "商品结算明细",
            ["订单号", "交易类型", "商品数量", "商品实付/实退", "佣金总额", "商品名称"],
            [order_id, "结算入账", 1, 214, -4.28, "手柄"],
        )
        order = workbook_bytes(
            "包裹详情",
            [
                "订单号", "订单状态", "售后状态", "商家编码", "商品名称",
                "SKU件数", "快递单号", "快递公司",
            ],
            [order_id, "已取消", "售后完成", "FF01A-04", "手柄", 1, "", ""],
        )
        raw = {
            "source_files": [
                {"platform": "小红书", "shop": "纷岚店", "fname": "小红书纷岚商品结算明细.xlsx", "buf": settlement},
                {"platform": "小红书", "shop": "纷岚店", "fname": "小红书纷岚订单明细.xlsx", "buf": order},
            ],
            "logistics": [],
        }

        result = settlement_engine.compute(
            raw,
            {"FF01A-04": {"unit_cost": 50, "source": "产品采购成本台", "name": "手柄"}},
            "2026-07",
        )

        self.assertTrue(any(
            row[1] == "小红书" and row[4] == "物流成本" and row[5] == order_id
            for row in result["gap_rows"]
        ))

    def test_confirmed_douyin_product_id_fills_blank_merchant_sku(self):
        order_id = "6953848598969259333"
        product_id = "3751006771863486516"
        settlement = (
            "订单号,结算单类型,收入合计,结算金额,商品数量,商品ID,商品名称,平台服务费\n"
            f"{order_id},已结算,214,204,1,{product_id},食人花2代,10\n"
        ).encode("utf-8-sig")
        order = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,"
            "订单提交时间,支付完成时间,订单状态,快递信息\n"
            f"{order_id},sub-1,食人花2代【磁吸+滑轨】,{product_id},,1,"
            "2026-07-10 10:00:00,2026-07-10 10:01:00,已完成,SF123-顺丰速运\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [
                {
                    "platform": "抖音", "shop": "宝空店",
                    "fname": "DL202608171903494679826433.csv", "buf": settlement,
                },
                {
                    "platform": "抖音", "shop": "宝空店",
                    "fname": "抖音宝空订单明细.csv", "buf": order,
                },
            ],
            "logistics": [
                {"tracking": "SF123", "carrier": "顺丰", "amount": 8, "source": "API"}
            ],
        }

        result = settlement_engine.compute(
            raw,
            {"PK02-S3": {"unit_cost": 100, "source": "产品采购成本台", "name": "食人花2代"}},
            "2026-07",
        )

        purchase_gaps = [
            row for row in result["gap_rows"]
            if row[1] == "抖音" and row[4] == "采购成本" and row[5] == order_id
        ]
        self.assertEqual([], purchase_gaps)
        matching_costs = [
            row for row in result["cost_rows"]
            if row[0] == "抖音" and row[3] == order_id
        ]
        self.assertEqual("PK02-S3", matching_costs[0][4])
        self.assertEqual(100, matching_costs[0][7])
        self.assertTrue(any(
            row[0:3] == ["抖音", "抖音宝空", "SKU映射确认"]
            and "3751006771863486516 → PK02-S3" in str(row[6])
            for row in result["source_rows"]
        ))

    def test_confirmed_no_settlement_is_structured_input_not_a_missing_file(self):
        raw = {
            "source_files": [
                {
                    "platform": "抖音",
                    "shop": "纷岚店",
                    "fname": "抖音纷岚订单.csv",
                    "buf": b"main,sub\n1,1\n",
                }
            ],
            "manifest_statuses": [
                {"platform": "抖音", "shop": "纷岚店", "file_type": "当月结算账单", "status": "已确认无数据"}
            ],
            "logistics": [],
        }

        result = settlement_engine.compute(raw, {}, "2026-07")

        rows = [row for row in result["monthly_rows"] if row[1:3] == ["抖音", "抖音纷岚"]]
        self.assertEqual(1, len(rows))
        self.assertEqual("运营已确认本月无结算，按0试算", rows[0][19])
        self.assertFalse(any(
            row[1:3] == ["抖音", "抖音纷岚"]
            and row[4] == "资料缺口"
            and row[5] == "抖音纷岚"
            for row in result["gap_rows"]
        ))
        self.assertTrue(any(
            row[1:3] == ["抖音", "抖音纷岚"]
            and row[4:6] == ["资料缺口", "广告账单"]
            for row in result["gap_rows"]
        ))

    def test_douyin_internal_entity_transfer_is_not_platform_fee(self):
        buf = (
            "动账时间,动账方向,动账金额,动账场景\n"
            "2026-07-01 10:00:00,出账,-500,主体变更旧主体资金转入新主体\n"
            "2026-07-02 10:00:00,出账,-0.58,权益保险\n"
        ).encode("utf-8-sig")

        rows = parsers.parse_dy_platform_fee(buf, "抖音纷岚平台费用.csv", "2026-07")

        self.assertEqual(1, len(rows))
        self.assertEqual("权益保险", rows[0]["fee_type"])
        self.assertEqual(-0.58, rows[0]["amount"])

    def test_douyin_public_other_fee_uses_sku_net_sales_and_reconciles_cents(self):
        settlement = (
            "订单号,结算单类型,收入合计,结算金额,商品数量,商品ID,商品名称,平台服务费\n"
            "order-a,已结算,100,100,1,product-a,商品A,0\n"
            "order-a,退款,-20,-20,0.2,product-a,商品A,0\n"
            "order-b,已结算,20,20,1,product-b,商品B,0\n"
        ).encode("utf-8-sig")
        orders = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,快递信息\n"
            "order-a,sub-a,商品A,product-a,SKU-A,1,SF-A-顺丰速运\n"
            "order-b,sub-b,商品B,product-b,SKU-B,1,SF-B-顺丰速运\n"
        ).encode("utf-8-sig")
        fees = (
            "动账时间,动账方向,动账金额,动账场景\n"
            "2026-07-31 10:00:00,出账,-1.01,权益保险\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音结算订单.csv", "buf": settlement},
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音宝空订单明细.csv", "buf": orders},
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音平台费用.csv", "buf": fees},
            ],
            "logistics": [
                {"tracking": "SF-A", "carrier": "顺丰", "amount": 0, "source": "账单"},
                {"tracking": "SF-B", "carrier": "顺丰", "amount": 0, "source": "账单"},
            ],
        }

        result = settlement_engine.compute(
            raw,
            {
                "SKU-A": {"unit_cost": 10, "source": "产品采购成本台", "name": "商品A"},
                "SKU-B": {"unit_cost": 10, "source": "产品采购成本台", "name": "商品B"},
            },
            "2026-07",
        )
        rows = {
            row[5]: row for row in result["product_rows"]
            if row[0] == "抖音" and row[3] == "抖音宝空"
        }

        self.assertAlmostEqual(rows["SKU-A"][15], 0.81, places=2)
        self.assertAlmostEqual(rows["SKU-B"][15], 0.20, places=2)
        self.assertAlmostEqual(sum(row[15] for row in rows.values()), 1.01, places=2)

    def test_finance_product_rows_use_erp_name_and_merge_platform_titles_by_sku(self):
        settlement = (
            "订单号,结算单类型,收入合计,结算金额,商品数量,商品ID,商品名称,平台服务费\n"
            "order-a,已结算,100,100,1,product-a,平台超长标题A,0\n"
            "order-b,已结算,50,50,1,product-b,平台超长标题B,0\n"
        ).encode("utf-8-sig")
        orders = (
            "主订单编号,子订单编号,选购商品,商品ID,商家编码,商品数量,快递信息\n"
            "order-a,sub-a,平台超长标题A,product-a,SKU-A,1,SF-A-顺丰速运\n"
            "order-b,sub-b,平台超长标题B,product-b,SKU-A,1,SF-B-顺丰速运\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音结算订单.csv", "buf": settlement},
                {"platform": "抖音", "shop": "宝空店", "fname": "抖音宝空订单明细.csv", "buf": orders},
            ],
            "logistics": [
                {"tracking": "SF-A", "carrier": "顺丰", "amount": 8, "source": "账单"},
                {"tracking": "SF-B", "carrier": "顺丰", "amount": 6, "source": "账单"},
            ],
        }

        result = settlement_engine.compute(
            raw,
            {"SKU-A": {"unit_cost": 10, "source": "产品采购成本台", "name": "ERP中文品名A"}},
            "2026-07",
        )
        rows = [
            row for row in result["product_rows"]
            if row[0] == "抖音" and row[3] == "抖音宝空" and row[5] == "SKU-A"
        ]

        self.assertEqual(1, len(rows))
        self.assertEqual("ERP中文品名A", rows[0][6])
        self.assertEqual(2, rows[0][7])
        self.assertEqual(150, rows[0][9])
        self.assertEqual(0, rows[0][11])
        self.assertEqual(20, rows[0][13])
        self.assertEqual(
            rows[0][9] - rows[0][10] - sum(rows[0][11:16]),
            rows[0][16],
        )
        self.assertTrue(all(row[5] == "ERP中文品名A" for row in result["cost_rows"]))

        missing_name_result = settlement_engine.compute(
            raw,
            {"SKU-A": {"unit_cost": 10, "source": "产品采购成本台", "name": ""}},
            "2026-07",
        )
        missing_name_rows = [
            row for row in missing_name_result["product_rows"]
            if row[0] == "抖音" and row[3] == "抖音宝空" and row[5] == "SKU-A"
        ]
        self.assertEqual("ERP品名未维护（SKU-A）", missing_name_rows[0][6])
        self.assertNotIn("平台超长标题", missing_name_rows[0][6])

    def test_powkong_tax_reporting_export_yields_one_precise_settlement_request(self):
        uploaded = (
            "业务时间,报送场景,报送金额,业务单号,明细数据对应菜单名称,用户实付,技术服务费\n"
            "2026-07-01,涉税报送,99,R001,资金报送,99,5\n"
        ).encode("utf-8-sig")
        raw = {
            "source_files": [{
                "platform": "抖音",
                "shop": "宝空店",
                "fname": "抖音 宝空 结算单.csv",
                "buf": uploaded,
            }],
            "logistics": [],
        }

        result = settlement_engine.compute(raw, {}, "2026-07")

        gaps = [
            row for row in result["gap_rows"]
            if row[1:3] == ["抖音", "抖音宝空"]
            and row[4] == "资料缺口"
            and row[5] != "广告账单"
        ]
        self.assertEqual(1, len(gaps))
        self.assertIn("已上传 抖音 宝空 结算单.csv", gaps[0][6])
        self.assertIn("订单号、结算单类型、收入合计、结算金额、商品数量、商品ID", gaps[0][6])
        self.assertIn("不要重复上传涉税报送明细", gaps[0][8])


class SettlementP0LedgerTests(unittest.IsolatedAsyncioTestCase):
    async def test_initial_check_only_never_creates_report_or_sends_new_card(self):
        raw = {
            "orders": [{"platform": "抖音", "shop": "抖音宝空", "sku": "PK02-S2", "tracking": ""}],
            "refunds": [],
            "plat_fees": [],
            "ads": [],
            "logistics": [],
            "sku_set": {"PK02-S2"},
            "source_files": [{"platform": "抖音", "shop": "抖音宝空", "fname": "抖音宝空结算单.csv", "buf": b"x"}],
            "skipped_shops": [],
            "shop_keys": {("抖音", "抖音宝空")},
        }
        settlement = {
            "gap_rows": [[
                "P0", "抖音", "抖音宝空", "2026-07", "资料缺口", "抖音宝空",
                "现有文件不是订单级结算明细", "无法计算", "补订单级结算明细",
            ]],
            "cost_rows": [],
        }
        refresh = {
            "active_message_ids": ["om_current"],
            "invalidated_message_ids": ["om_old"],
            "missing_existing_cards": 0,
        }
        with patch("builtins.print"), patch.object(task_runner, "update_status", new=AsyncMock()) as update_status, patch.object(
            task_runner, "get_record", new=AsyncMock(return_value={"月份": "2026-07"})
        ), patch.object(
            task_runner, "collect_raw_data", new=AsyncMock(return_value=raw)
        ), patch.object(
            task_runner.lingxing, "get_products", new=AsyncMock(return_value={})
        ), patch.object(
            task_runner, "_load_finance_cost_map", new=AsyncMock(return_value={})
        ), patch.object(
            task_runner.config, "SF_API_ENABLED", False
        ), patch.object(
            task_runner.engine, "compute", return_value={"shop_totals": {}}
        ), patch.object(
            task_runner.settlement_engine, "compute", return_value=settlement
        ), patch.object(
            task_runner, "sync_settlement_p0_gaps", new=AsyncMock(return_value={})
        ), patch.object(
            task_runner.cost_gap_alert, "refresh_existing_operation_gap_cards", new=AsyncMock(return_value=refresh)
        ) as refresh_cards, patch.object(
            task_runner.ledger, "open_p0_gaps", new=AsyncMock(return_value=[{"fields": {"gap_id": "gap_1"}}])
        ), patch.object(
            task_runner.ledger, "update_run", new=AsyncMock()
        ), patch.object(
            task_runner.writer, "create_report_spreadsheet", new=AsyncMock(side_effect=AssertionError("不得生成报表"))
        ) as create_report, patch.object(
            task_runner, "_notify_report_ready", new=AsyncMock(side_effect=AssertionError("不得发新消息"))
        ) as notify:
            result = await task_runner.run_profit(
                "rec_2026_07",
                suppress_notify=True,
                initial_check_only=True,
            )

        self.assertTrue(result["ok"])
        self.assertFalse(result["created_report"])
        self.assertFalse(result["sent_new_card"])
        refresh_cards.assert_awaited_once()
        update_status.assert_not_awaited()
        create_report.assert_not_awaited()
        notify.assert_not_awaited()

    async def test_calculated_p0_uses_existing_base_select_options(self):
        settlement = {
            "gap_rows": [
                ["P0", "抖音", "抖音宝空", "2026-07", "资料缺口", "抖音宝空", "缺结算明细", "不能核算", "补结算明细"],
                ["P0", "天猫", "天猫宝空", "2026-07", "物流成本", "SF123", "缺物流费", "毛利偏高", "补物流账单"],
            ]
        }

        async def create_gap(_run_id, gap_type, _platform, _month, evidence):
            self.assertIn(gap_type, {"其他", "ERP_SKU缺失", "物流账单缺失"})
            self.assertTrue(evidence.startswith("[初检计算]"))
            return {"fields": {"gap_id": f"gap-{gap_type}"}}

        with patch.object(
            task_runner.ledger,
            "create_gap",
            new=AsyncMock(side_effect=create_gap),
        ) as create, patch.object(
            task_runner.ledger,
            "mark_gap",
            new=AsyncMock(),
        ), patch.object(
            task_runner.ledger,
            "gaps_for_run",
            new=AsyncMock(return_value=[]),
        ):
            result = await task_runner.sync_settlement_p0_gaps("2026-07", settlement)

        self.assertEqual(2, result["open_calculated_p0"])
        self.assertEqual("其他", create.await_args_list[0].args[1])
        self.assertEqual("物流账单缺失", create.await_args_list[1].args[1])



class DouyinP0RegressionTests(unittest.TestCase):
    def test_douyin_finance_flow_uses_cash_spend_and_keeps_gifts_as_memo(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "日期", "余额总消耗(元)", "非赠款消耗(元)", "赠款消耗(元)",
            "消返红包消耗(元)", "立减红包消耗(元)", "共享钱包消耗(元)",
            "共享赠款消耗(元)", "总存入(元)",
        ])
        sheet.append(["总计", 22407.48, 22389.99, 17.49, 200, 0, 0, 0, 22701.49])
        buf = io.BytesIO()
        workbook.save(buf)

        parsed = parsers.detect_and_parse(
            "财务流水.xlsx",
            buf.getvalue(),
            "2026-08",
            "广告",
            platform="抖音",
        )
        rows = parsed["data"]

        self.assertEqual(1, len(rows))
        self.assertAlmostEqual(22389.99, rows[0]["spend"], places=2)
        self.assertAlmostEqual(17.49, rows[0]["gift_spend"], places=2)
        self.assertAlmostEqual(200.00, rows[0]["red_packet_spend"], places=2)
        self.assertAlmostEqual(22701.49, rows[0]["deposit"], places=2)
        self.assertEqual("财务流水", rows[0]["source_type"])

    def test_douyin_confirmed_expense_policy_reconciles_profit_and_settlement_payback(self):
        settlement = _csv_bytes(
            ["订单号", "结算单类型", "收入合计", "结算金额", "商品ID", "商品数量", "平台服务费", "达人佣金", "站外推广费"],
            [["10001", "已结算", 100, 75, "P1", 1, 5, 15, 5]],
        )
        orders = _csv_bytes(
            ["主订单编号", "商品ID", "商家编码", "商品数量", "订单状态", "售后状态", "快递信息"],
            [["10001", "P1", "SKU-A", 1, "已完成", "", "SF10001-顺丰速运"]],
        )
        account = _csv_bytes(
            ["动账时间", "动账方向", "动账场景", "动账金额", "备注"],
            [
                ["2026-08-01", "出账", "退换货运费险", 2, ""],
                ["2026-08-01", "出账", "消费者赔付", 3, ""],
                ["2026-08-01", "出账", "抖音月付联合贴息", 4, ""],
                ["2026-08-01", "出账", "小额打款", 40, "补差价"],
                ["2026-08-01", "出账", "评价有礼", 15, ""],
                ["2026-08-01", "出账", "巨量千川充值", 100, ""],
                ["2026-08-01", "出账", "提现", 50, ""],
            ],
        )
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "结算明细.csv", "buf": settlement, "attach_field": "结算明细"},
                {"platform": "抖音", "shop": "宝空店", "fname": "订单明细.csv", "buf": orders, "attach_field": "订单明细"},
                {"platform": "抖音", "shop": "宝空店", "fname": "动账明细.csv", "buf": account, "attach_field": "动账明细"},
            ],
            "ads": [
                {
                    "platform": "抖音",
                    "shop": "宝空店",
                    "source_type": "财务流水",
                    "spend": 10,
                    "gift_spend": 2,
                    "red_packet_spend": 3,
                    "source": "财务流水.xlsx",
                }
            ],
            "logistics": [{"tracking": "SF10001", "amount": 8, "source": "顺丰账单"}],
        }

        report = settlement_engine.compute(
            raw,
            {"SKU-A": {"unit_cost": 10, "name": "测试产品", "source": "成本台"}},
            "2026-08",
        )
        monthly = next(row for row in report["monthly_rows"] if row[1:3] == ["抖音", "抖音宝空"])

        self.assertAlmostEqual(100.00, monthly[6], places=2)
        self.assertAlmostEqual(40.00, monthly[7], places=2)
        self.assertAlmostEqual(60.00, monthly[8], places=2)
        self.assertAlmostEqual(25.00, monthly[9], places=2)
        self.assertAlmostEqual(10.00, monthly[10], places=2)
        self.assertAlmostEqual(10.00, monthly[11], places=2)
        self.assertAlmostEqual(8.00, monthly[12], places=2)
        self.assertAlmostEqual(9.00, monthly[13], places=2)
        self.assertAlmostEqual(-2.00, monthly[14], places=2)
        self.assertAlmostEqual(75.00, monthly[16], places=2)
        self.assertFalse([row for row in report["gap_rows"] if row[0] == "P0" and row[2] == "抖音宝空"])
        self.assertEqual(
            [["抖音", "抖音宝空", "2026-08", "动账明细.csv",
              "其他流动资产-抖店评价有礼活动保证金", "", 15.0, "动账金额",
              "出账：保证金性质，不计当期毛利费用"]],
            report["asset_rows"],
        )

    def test_douyin_same_product_id_multi_sku_order_costs_each_settlement_line(self):
        settlement = _csv_bytes(
            ["订单号", "结算单类型", "收入合计", "结算金额", "商品ID", "商品数量", "平台服务费", "达人佣金", "站外推广费"],
            [
                ["6928676009172237729", "已结算", 381.09, 316.31, "P1", 1, 7.62, 57.16, 0],
                ["6928676009172237729", "已结算", 256.91, 213.23, "P1", 1, 5.14, 38.54, 0],
            ],
        )
        orders = _csv_bytes(
            ["主订单编号", "商品ID", "商家编码", "商品数量", "订单状态", "售后状态", "快递信息"],
            [
                ["6928676009172237729", "P1", "PK02-S2", 1, "已完成", "", "76967441851947-顺丰速运"],
                ["6928676009172237729", "P1", "PK01pro", 1, "已完成", "", "76967471561416-顺丰速运"],
            ],
        )
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "结算明细.csv", "buf": settlement, "attach_field": "结算明细"},
                {"platform": "抖音", "shop": "宝空店", "fname": "订单明细.csv", "buf": orders, "attach_field": "订单明细"},
            ],
            "ads": [],
            "logistics": [
                {"tracking": "76967441851947", "amount": 5, "source": "顺丰账单"},
                {"tracking": "76967471561416", "amount": 6, "source": "顺丰账单"},
            ],
        }

        report = settlement_engine.compute(
            raw,
            {
                "PK02-S2": {"unit_cost": 173.73, "name": "产品A", "source": "成本台"},
                "PK01pro": {"unit_cost": 100.79, "name": "产品B", "source": "成本台"},
            },
            "2026-08",
        )

        self.assertAlmostEqual(274.52, next(row for row in report["monthly_rows"] if row[1:3] == ["抖音", "抖音宝空"])[11], places=2)
        self.assertEqual({"PK02-S2", "PK01pro"}, {row[4] for row in report["cost_rows"]})
        self.assertFalse([row for row in report["gap_rows"] if row[4] == "采购成本"])

    def test_douyin_unshipped_refund_has_no_purchase_cost_or_freight_gap(self):
        settlement = _csv_bytes(
            ["订单号", "结算单类型", "收入合计", "结算金额", "商品ID", "商品数量", "平台服务费", "达人佣金", "站外推广费"],
            [["6955468382802745067", "已结算", 1, 1, "P1", 1, 0, 0, 0]],
        )
        orders = _csv_bytes(
            ["主订单编号", "商品ID", "商家编码", "商品数量", "订单状态", "售后状态", "快递信息"],
            [["6955468382802745067", "P1", "PK02-S3", 1, "已关闭", "退款成功", "-"]],
        )
        raw = {
            "source_files": [
                {"platform": "抖音", "shop": "宝空店", "fname": "结算明细.csv", "buf": settlement, "attach_field": "结算明细"},
                {"platform": "抖音", "shop": "宝空店", "fname": "订单明细.csv", "buf": orders, "attach_field": "订单明细"},
            ],
            "ads": [],
            "logistics": [],
        }

        report = settlement_engine.compute(
            raw,
            {"PK02-S3": {"unit_cost": 173.73, "name": "产品C", "source": "成本台"}},
            "2026-08",
        )

        self.assertAlmostEqual(0.00, next(row for row in report["monthly_rows"] if row[1:3] == ["抖音", "抖音宝空"])[11], places=2)
        self.assertAlmostEqual(0.00, next(row for row in report["monthly_rows"] if row[1:3] == ["抖音", "抖音宝空"])[12], places=2)
        self.assertFalse([row for row in report["gap_rows"] if row[4] == "物流成本"])
        self.assertTrue(any(row[3] == "6955468382802745067" and row[11] == "无需运费" for row in report["log_rows"]))

if __name__ == "__main__":
    unittest.main()
